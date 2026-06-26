"""Tests for the spec->DTA generator and renderer (scripts/dataset_spec_to_dta.py,
scripts/render_dta_docx.py).

Covers: format auto-detection, CSV type inference, domain-selection rules,
the "never invent data" invariants, ODRL policy and multi-standard mappings,
the Word render, and a linkml-gated schema-conformance check.
"""
from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parent.parent
SCRIPTS = REPO / "scripts"
SCHEMA = REPO / "dds.yaml"


def _load(module_name: str):
    spec = importlib.util.spec_from_file_location(module_name, SCRIPTS / f"{module_name}.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = mod
    spec.loader.exec_module(mod)
    return mod


dta = _load("dataset_spec_to_dta")
render = _load("render_dta_docx")


# --- fixtures -----------------------------------------------------------------

def define_spec() -> dict:
    return {
        "studyName": "S1", "protocolName": "P1", "fileType": "Data Definition Specification",
        "itemGroups": [
            {"OID": "IG.LB", "name": "LB", "domain": "LB",
             "structure": "one per result",
             "items": [{"OID": "IT.LB.X", "name": "X", "dataType": "float"}]},
            {"OID": "IG.VS", "name": "VS", "domain": "VS", "items": []},
        ],
    }


def single_group_spec() -> dict:
    return {"itemGroups": [{"OID": "IG.ONE", "name": "ONE", "domain": "ONE",
                            "items": [{"OID": "IT.A", "name": "A", "dataType": "text"}]}]}


# --- format detection ---------------------------------------------------------

def test_load_data_definition_spec_passthrough(tmp_path):
    p = tmp_path / "d.json"
    p.write_text(json.dumps(define_spec()))
    assert "itemGroups" in dta.load_spec(p)


def test_load_dataset_json_maps_columns_and_normalises_dtype(tmp_path):
    dsj = {"clinicalData": {"studyOID": "STDY", "itemGroupData": {
        "IG.LB": {"columns": [
            {"name": "USUBJID", "label": "Subject", "dataType": "string"},
            {"name": "VAL", "label": "Value", "dataType": "float", "length": 5},
        ]}}}}
    p = tmp_path / "ds.json"
    p.write_text(json.dumps(dsj))
    spec = dta.load_spec(p)
    cols = spec["itemGroups"][0]["items"]
    assert spec["itemGroups"][0]["domain"] == "LB"
    assert cols[0]["dataType"] == "text"   # "string" -> "text"
    assert cols[1]["dataType"] == "float"


def test_load_csv_infers_types(tmp_path):
    p = tmp_path / "lab.csv"
    p.write_text("ID,AGE,VAL,DT,NOTE\n1,40,3.5,2024-01-15,hello\n2,41,4.0,2024-02-15,world\n")
    items = {i["name"]: i["dataType"] for i in dta.load_spec(p)["itemGroups"][0]["items"]}
    assert items == {"ID": "integer", "AGE": "integer", "VAL": "float",
                     "DT": "datetime", "NOTE": "text"}


def test_load_yaml_requires_itemgroups(tmp_path):
    p = tmp_path / "x.yaml"
    p.write_text("foo: bar\n")
    with pytest.raises(SystemExit):
        dta.load_spec(p)


def test_xml_routed_to_converter(tmp_path):
    p = tmp_path / "d.xml"
    p.write_text("<ODM/>")
    with pytest.raises(SystemExit) as e:
        dta.load_spec(p)
    assert "xml2json" in str(e.value)


def test_unsupported_and_malformed(tmp_path):
    with pytest.raises(SystemExit):
        dta.load_spec(tmp_path / "missing.json")
    bad = tmp_path / "b.json"
    bad.write_text("{not json")
    with pytest.raises(SystemExit):
        dta.load_spec(bad)
    weird = tmp_path / "w.json"
    weird.write_text(json.dumps({"hello": 1}))
    with pytest.raises(SystemExit):
        dta.load_spec(weird)


# --- domain selection ---------------------------------------------------------

def test_domain_optional_for_single_group():
    g = dta.find_item_group(single_group_spec(), None)
    assert g["domain"] == "ONE"


def test_domain_required_when_multiple():
    with pytest.raises(SystemExit) as e:
        dta.find_item_group(define_spec(), None)
    assert "LB" in str(e.value) and "VS" in str(e.value)


def test_unknown_domain_lists_available():
    with pytest.raises(SystemExit) as e:
        dta.find_item_group(define_spec(), "ZZ")
    assert "Available" in str(e.value)


# --- "never invent data" invariants ------------------------------------------

def test_unsupplied_party_and_terms_are_placeholders():
    spec = define_spec()
    out = dta.build_dta(spec, spec["itemGroups"][0],
                        provider_name=dta.PLACEHOLDER, provider_type="Lab",
                        consumer_name=dta.PLACEHOLDER)
    assert out["provider"]["name"] == dta.PLACEHOLDER
    assert out["consumer"] == dta.PLACEHOLDER
    assert "hasPolicy" not in out                      # no purpose -> no policy
    assert "deliverySchedule" not in out["dataFlow"]   # no cadence -> no schedule


def test_golden_structure_for_define_input():
    spec = define_spec()
    out = dta.build_dta(spec, spec["itemGroups"][0],
                        provider_name="ACME Labs", provider_type="Lab",
                        consumer_name="Sponsor Inc")
    assert out["OID"] == "PA.DTA.LB"
    assert out["provider"]["name"] == "ACME Labs"
    assert out["provider"]["type"] == "Lab"
    dsd = out["dataFlow"]["structure"]
    assert dsd["domain"] == "LB"
    assert [i["name"] for i in dsd["items"]] == ["X"]


def test_cadence_is_neutral_string():
    spec = single_group_spec()
    out = dta.build_dta(spec, spec["itemGroups"][0],
                        provider_name="p", provider_type="Lab", consumer_name="c",
                        cadence="R/2025-01-01/P1M")
    assert out["dataFlow"]["deliverySchedule"] == ["R/2025-01-01/P1M"]


def test_permitted_purpose_builds_odrl_policy():
    spec = single_group_spec()
    out = dta.build_dta(spec, spec["itemGroups"][0],
                        provider_name="p", provider_type="Lab", consumer_name="c",
                        permitted_purpose="safety-reporting")
    rule = out["hasPolicy"][0]["permission"][0]
    assert rule["action"] == "use"
    assert rule["constraint"][0]["rightOperand"] == "safety-reporting"
    assert rule["constraint"][0]["operator"] == "eq"


def test_mappings_merge_codings_and_concept():
    spec = single_group_spec()
    maps = {"implementsConcept": "RC.X",
            "variables": {"A": {"coding": [{"codeSystem": "LOINC", "code": "1-1"}]}}}
    out = dta.build_dta(spec, spec["itemGroups"][0],
                        provider_name="p", provider_type="Lab", consumer_name="c",
                        mappings=maps)
    dsd = out["dataFlow"]["structure"]
    assert dsd["implementsConcept"] == "RC.X"
    assert dsd["items"][0]["coding"][0]["code"] == "1-1"


# --- render -------------------------------------------------------------------

def test_render_docx_has_sections_and_table(tmp_path):
    pytest.importorskip("docx")
    spec = single_group_spec()
    out = dta.build_dta(spec, spec["itemGroups"][0],
                        provider_name="p", provider_type="Lab", consumer_name="c",
                        cadence="R/2025-01-01/P1M", permitted_purpose="x")
    doc = render.build_doc(out, "src.json")
    f = tmp_path / "o.docx"
    doc.save(str(f))
    from docx import Document
    d = Document(str(f))
    headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert any("Variable specification" in h for h in headings)
    assert any("Governance" in h for h in headings)


# --- schema conformance (gated on linkml + schema presence) -------------------

@pytest.mark.skipif(shutil.which("linkml-validate") is None or not SCHEMA.exists(),
                    reason="linkml-validate or dds.yaml not available")
def test_generated_dta_is_schema_valid(tmp_path):
    spec = define_spec()
    out = dta.build_dta(spec, spec["itemGroups"][0],
                        provider_name="ACME", provider_type="Lab", consumer_name="S",
                        cadence="R/2025-01-01/P1M", permitted_purpose="safety-reporting")
    f = tmp_path / "dta.json"
    f.write_text(json.dumps(out))
    r = subprocess.run(["linkml-validate", "-s", str(SCHEMA),
                        "-C", "ProvisionAgreement", str(f)],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr
